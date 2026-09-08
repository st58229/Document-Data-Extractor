import math
import tkinter as tk
from tkinter import filedialog, ttk, messagebox
import pandas as pd
import os
import json
import fitz  # PyMuPDF
from docx import Document
from datetime import datetime
import re
from openpyxl import load_workbook
import time
import csv
import psutil


# Initialization of the main window
root = tk.Tk()
root.title("Big Data Processing App")
root.geometry("500x300")

# Global variables
files_set = set()
processed_data = {}
excel_extraction_settings = []
word_extraction_settings = []

# Global settings
settings = {
    'delimiter': ';',
    'decimal': ','
}

# Function to select a folder and files within it
def select_folder():
    folder_path = filedialog.askdirectory()
    if folder_path:
        for root_dir, _, files in os.walk(folder_path):
            for file in files:
                file_path = os.path.join(root_dir, file).replace("\\", "/")
                files_set.add(file_path)

# Function to select files
def select_files():
    file_paths = filedialog.askopenfilenames()
    if file_paths:
        files_set.update(file_paths)

# Display file information
def show_selected_files_info():
    if not files_set:
        messagebox.showinfo("Info", "No files were selected.")
        return

    total_size = convert_size(sum(os.path.getsize(file) for file in files_set))
    message = (
        f"Number of selected files: {len(files_set)}\n"
        f"Total size: {total_size}\n\n"
        "Selected files:\n" + "\n".join(files_set)
    )

    info_window = tk.Toplevel()
    info_window.title("File Information")
    info_window.geometry("600x400")

    text_box = tk.Text(info_window, wrap='word')
    scroll_bar = tk.Scrollbar(info_window, command=text_box.yview)
    text_box.config(yscrollcommand=scroll_bar.set)

    text_box.insert('1.0', message)
    text_box.config(state='disabled')

    text_box.pack(side='left', fill='both', expand=True)
    scroll_bar.pack(side='right', fill='y')

# Convert file size
def convert_size(size_bytes):
    if size_bytes == 0:
        return "0B"
    size_name = ("B", "kB", "MB", "GB", "TB")
    i = int(math.floor(math.log(size_bytes, 1024)))
    p = math.pow(1024, i)
    s = round(size_bytes / p, 2)
    return f"{s} {size_name[i]}"

# Function to format time
def format_elapsed_time(elapsed_time):
    if elapsed_time < 60:
        return f"{int(elapsed_time)} seconds"
    elif elapsed_time < 3600:
        minutes = int(elapsed_time // 60)
        seconds = int(elapsed_time % 60)
        return f"{minutes} minutes a {seconds} seconds"
    else:
        hours = int(elapsed_time // 3600)
        minutes = int((elapsed_time % 3600) // 60)
        return f"{hours} hours a {minutes} minutes"

# JSON loading and export
def load_json(file_path):
    with open(file_path, 'r', encoding='utf-8') as file:
        return json.load(file)

def export_to_json(data):
    json_file_path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON files", "*.json")])
    if not json_file_path:
        messagebox.showinfo("Info", "Export was canceled.")
        return

    with open(json_file_path, 'w', encoding='utf-8') as json_file:
        json.dump(data, json_file, ensure_ascii=False, indent=4)

    messagebox.showinfo("Success", "Data were exported to a JSON file.")

# Function for settings
def open_settings():
    settings_window = tk.Toplevel(root)
    settings_window.title("Delimiter Settings")
    settings_window.geometry("300x200")

    delimiter_label = tk.Label(settings_window, text="CSV Delimiter:")
    delimiter_label.pack(pady=5)

    delimiter_entry = tk.Entry(settings_window)
    delimiter_entry.insert(0, settings['delimiter'])
    delimiter_entry.pack(pady=5)

    decimal_label = tk.Label(settings_window, text="CSV Decimal:")
    decimal_label.pack(pady=5)

    decimal_entry = tk.Entry(settings_window)
    decimal_entry.insert(0, settings['decimal'])
    decimal_entry.pack(pady=5)

    def save_settings():
        settings['delimiter'] = delimiter_entry.get() or ';'
        settings['decimal'] = decimal_entry.get() or ','
        settings_window.destroy()

    save_button = tk.Button(settings_window, text="Save", command=save_settings)
    save_button.pack(pady=10)

# Function to export list of extracted cells
def export_extraction_settings(extraction_settings):
    file_path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON files", "*.json")])
    if not file_path:
        return
    with open(file_path, 'w', encoding='utf-8') as file:
        json.dump(extraction_settings, file, ensure_ascii=False, indent=4)
    messagebox.showinfo("Export", "Extraction settings were successfully exported.")

# Function to import list of extracted cells
def import_extraction_settings(extraction_list, extraction_settings):
    file_path = filedialog.askopenfilename(filetypes=[("JSON files", "*.json")])
    if not file_path:
        return
    with open(file_path, 'r', encoding='utf-8') as file:
        imported_settings = json.load(file)
    extraction_settings.clear()
    extraction_settings.extend(imported_settings)
    extraction_list.delete(0, tk.END)
    for setting in extraction_settings:
        if extraction_settings == excel_extraction_settings:
            extraction_list.insert(tk.END, f"{setting['cell']}: {setting['attribute']}")
        else:
            extraction_list.insert(tk.END, f"{setting['keyword']}: {setting['attribute']}")
    messagebox.showinfo("Import", "Extraction settings were successfully imported.")

# Function for Excel extraction settings
def open_excel_extraction_settings():
    def save_excel_extraction():
        cell = cell_entry.get()
        attribute = attribute_entry.get()
        if not cell or not attribute:
            messagebox.showwarning("Warning", "You must enter both cell and attribute name.")
            return

        if not re.match(r'^[A-Za-z]+\d+$', cell):
            messagebox.showerror("Error", "Cell must be in valid format, e.g. A1 or B2.")
            return

        excel_extraction_settings.append({'cell': cell, 'attribute': attribute})
        extraction_list.insert(tk.END, f"{cell}: {attribute}")

        cell_entry.delete(0, tk.END)
        attribute_entry.delete(0, tk.END)

    def load_existing_settings():
        extraction_list.delete(0, tk.END)
        for setting in excel_extraction_settings:
            extraction_list.insert(tk.END, f"{setting['cell']}: {setting['attribute']}")

    extraction_window = tk.Toplevel(root)
    extraction_window.title("Excel Extraction Settings")
    extraction_window.geometry("400x450")

    cell_label = tk.Label(extraction_window, text="Cell:")
    cell_label.pack(pady=5)

    cell_entry = tk.Entry(extraction_window)
    cell_entry.pack(pady=5)

    attribute_label = tk.Label(extraction_window, text="Attribute name:")
    attribute_label.pack(pady=5)

    attribute_entry = tk.Entry(extraction_window)
    attribute_entry.pack(pady=5)

    save_button = tk.Button(extraction_window, text="Save", command=save_excel_extraction)
    save_button.pack(pady=10)

    extraction_list = tk.Listbox(extraction_window)
    extraction_list.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    export_button = tk.Button(extraction_window, text="Export list", command=lambda: export_extraction_settings(excel_extraction_settings))
    export_button.pack(pady=5)

    import_button = tk.Button(extraction_window, text="Import list", command=lambda: import_extraction_settings(extraction_list, excel_extraction_settings))
    import_button.pack(pady=5)

    load_existing_settings()

def open_word_extraction_settings():
    def save_word_extraction():
        keyword = keyword_entry.get()
        attribute = attribute_entry.get()
        if not keyword or not attribute:
            messagebox.showwarning("Warning", "You must enter both keyword and attribute name.")
            return

        word_extraction_settings.append({'keyword': keyword, 'attribute': attribute})
        extraction_list.insert(tk.END, f"{keyword}: {attribute}")

        keyword_entry.delete(0, tk.END)
        attribute_entry.delete(0, tk.END)

    def load_existing_settings():
        extraction_list.delete(0, tk.END)
        for setting in word_extraction_settings:
            extraction_list.insert(tk.END, f"{setting['keyword']}: {setting['attribute']}")

    extraction_window = tk.Toplevel(root)
    extraction_window.title("Word and PDF Extraction Settings")
    extraction_window.geometry("400x450")

    keyword_label = tk.Label(extraction_window, text="Keyword:")
    keyword_label.pack(pady=5)

    keyword_entry = tk.Entry(extraction_window)
    keyword_entry.pack(pady=5)

    attribute_label = tk.Label(extraction_window, text="Attribute name:")
    attribute_label.pack(pady=5)

    attribute_entry = tk.Entry(extraction_window)
    attribute_entry.pack(pady=5)

    save_button = tk.Button(extraction_window, text="Save", command=save_word_extraction)
    save_button.pack(pady=10)

    extraction_list = tk.Listbox(extraction_window)
    extraction_list.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    export_button = tk.Button(extraction_window, text="Export list", command=lambda: export_extraction_settings(word_extraction_settings))
    export_button.pack(pady=5)

    import_button = tk.Button(extraction_window, text="Import list", command=lambda: import_extraction_settings(extraction_list, word_extraction_settings))
    import_button.pack(pady=5)

    load_existing_settings()

# Tool for processing structured CSV and Excel files
def open_structured_data_tool():

        structured_file_path = filedialog.askopenfilename(filetypes=[("CSV or Excel files", "*.csv *.xls *.xlsx")])
        if not structured_file_path:
            return

        if structured_file_path.endswith('.csv'):
            df = pd.read_csv(structured_file_path, delimiter=settings['delimiter'], decimal=settings['decimal'])
        else:
            df = pd.read_excel(structured_file_path)

        json_data = df.to_dict(orient='records')
        export_to_json(json_data)

# Process Excel files
def process_excel(file_path):
    wb = load_workbook(file_path, data_only=True)
    sheet = wb.active

    if not excel_extraction_settings:
        messagebox.showerror("Error", "No cells were specified for extraction.")
        return {}

    extracted_data = {}
    for setting in excel_extraction_settings:
        cell = setting['cell']
        attribute = setting['attribute']
        try:
            value = sheet[cell].value
            extracted_data[attribute] = value
        except Exception:
            extracted_data[attribute] = None

    return extracted_data

# Process PDF and DOC files
def process_doc(file_path):
    with fitz.open(file_path) as doc:
        text = "".join(page.get_text() for page in doc)

    return extract_data_by_keywords(text)

# Process files
def process_files(file_type):
    global processed_data
    if not files_set:
        messagebox.showerror("Error", "Please select files first.")
        return

    process = psutil.Process(os.getpid())

    start_time = time.time()
    progress_bar['value'] = 0
    progress_bar['maximum'] = len(files_set)

    processed_data = {}
    skipped_files = []
    for idx, file_path in enumerate(files_set):
        ext = os.path.splitext(file_path)[1].lower()
        try:
            # Same functions, different data. Requires input for distinction
            if ext in ['.xls', '.xlsx'] and file_type == "excel":
                processed_data[file_path] = process_excel(file_path)
            elif ext == '.pdf' and file_type == "wrd/pdf":
                processed_data[file_path] = process_doc(file_path)
            elif ext == '.docx' and file_type == "wrd/pdf":
                processed_data[file_path] = process_doc(file_path)
            else:
                skipped_files.append(file_path)
        except Exception as e:
            skipped_files.append(f"{file_path} (Error: {str(e)})")
        finally:
            progress_bar['value'] = idx + 1
            root.update_idletasks()

    end_time = time.time()
    peak_mem_mb = process.memory_info().peak_wset / (1024 * 1024)

    elapsed_time = end_time - start_time

    elapsed_time_str = format_elapsed_time(elapsed_time)

    if skipped_files:
        skipped_message = "\n".join(skipped_files)
        log_unprocessed_files(skipped_files)
        messagebox.showwarning("Warning", f"Some files were not processed:\n{skipped_message}")

    files_set.clear()
    messagebox.showinfo("Success", f"File processing completed.\nElapsed time: {elapsed_time_str}\nMaximum RAM: {peak_mem_mb:.2f} MB")


def merge_json_files():
    # Select JSON files
    file_paths = filedialog.askopenfilenames(filetypes=[("JSON Files", "*.json")])
    if not file_paths:
        messagebox.showinfo("Info", "No file was selected.")
        return

    merged_data = {}

    for file_path in file_paths:
        try:
            with open(file_path, "r", encoding="utf-8") as file:
                data = json.load(file)

                if isinstance(data, dict):
                    merged_data.update(data)  # Direct merge without nesting
                else:
                    messagebox.showwarning("Warning", f"The file {file_path} contains an unsupported structure.")
        except Exception as e:
            messagebox.showerror("Error", f"Error reading the file {file_path}: {e}")

    # Select path to save the resulting JSON
    save_path = filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON Files", "*.json")])
    if save_path:
        with open(save_path, "w", encoding="utf-8") as output_file:
            json.dump(merged_data, output_file, ensure_ascii=False, indent=4)
        messagebox.showinfo("Success", f"JSON successfully merged and saved to {save_path}")

# Log unprocessed files
def log_unprocessed_files(skipped_files):
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    log_filename = f"Unprocessed_files_{timestamp}.txt"
    with open(log_filename, 'w', encoding='utf-8') as log_file:
        log_file.write("\n".join(skipped_files))

# Application information function
def about_app():
    messagebox.showinfo("About", "Big Data Processing App\nVersion: 1.0\nAuthor: Tomáš Fráňa\nst58229@upce.cz")

def extract_text_from_pdf(pdf_path):
    text_data = []
    try:
        doc = fitz.open(pdf_path)
        for page in doc:
            text_data.append(page.get_text())
        return "\n".join(text_data)
    except Exception as e:
        print(f"Error processing PDF: {e}")
        return None

def extract_text_from_docx(docx_path):
    try:
        doc = Document(docx_path)
        return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    except Exception as e:
        print(f"Error processing Word file: {e}")
        return None

def extract_data_by_keywords(text):
    extracted_data = {}
    for setting in word_extraction_settings:
        pattern = rf"{setting['keyword']}\s*[:\-]?\s*(.+)"
        attribute = setting['attribute']
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            extracted_data[attribute] = match.group(1).strip()
    return extracted_data

def json_to_csv():
    json_file = filedialog.askopenfilename(filetypes=[("JSON Files", "*.json")])
    if not json_file:
        messagebox.showinfo("Info", "No file was selected.")
        return

    csv_file = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV Files", "*.csv")])
    if not csv_file:
        return

    try:
        # Load JSON file
        with open(json_file, "r", encoding="utf-8") as file:
            data = json.load(file)

        if not isinstance(data, dict):  # Ensure JSON is a dict (not a list)
            messagebox.showerror("Error", "JSON file must be a dictionary where keys are file paths.")
            return

        structured_data = []
        header = set(["file_path"])  # Header includes file_path

        # Convert JSON to list of dicts
        for file_path, attributes in data.items():
            if isinstance(attributes, dict):  # Ensure values are dicts
                header.update(attributes.keys())  # Add keys to header
                record = {"file_path": file_path, **attributes}  # Add path to the file
                structured_data.append(record)

        # Order keys for consistency
        header = sorted(header)

        # Write to CSV
        with open(csv_file, "w", encoding="utf-8-sig", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=header)
            writer.writeheader()
            writer.writerows(structured_data)

        messagebox.showinfo("Success", f"JSON was successfully converted and saved as CSV: {csv_file}")

    except Exception as e:
        messagebox.showerror("Error", f"Error converting JSON to CSV: {e}")

# GUI elements
menu_bar = tk.Menu(root)
root.config(menu=menu_bar)

csv_menu = tk.Menu(menu_bar, tearoff=0)
menu_bar.add_cascade(label="CSV -> JSON", menu=csv_menu)
csv_menu.add_command(label="Convert CSV", command=open_structured_data_tool)
csv_menu.add_command(label="Delimiter Settings", command=open_settings)

menu_bar.add_command(label="Merge JSON", command=merge_json_files)
menu_bar.add_command(label="JSON -> CSV", command=json_to_csv)
menu_bar.add_command(label="About", command=about_app)

frame_select = tk.Frame(root)
frame_select.pack(pady=5)

file_button = tk.Button(frame_select, text="Select Files", command=select_files)
folder_button = tk.Button(frame_select, text="Select Folder", command=select_folder)
info_button = tk.Button(frame_select, text="ℹ", command=show_selected_files_info, width=3)

file_button.pack(side="left", padx=2)
folder_button.pack(side="left", padx=2)
info_button.pack(side="left", padx=2)

frame_main = tk.Frame(root)
frame_main.pack(pady=10, padx=10, fill="both", expand=True)

frame_left = tk.Frame(frame_main)
frame_left.pack(side="left", padx=10, fill="both", expand=True)

frame_right = tk.Frame(frame_main)
frame_right.pack(side="right", padx=10, fill="both", expand=True)

# Buttons for Excel processing
btn_settings_excel = tk.Button(frame_left, text="Excel Extraction Settings", command=open_excel_extraction_settings)
btn_settings_excel.pack(pady=5, fill="x")

btn_process_excel = tk.Button(frame_left, text="Process Files (Excel)", command=lambda: process_files("excel"))
btn_process_excel.pack(pady=5, fill="x")

# Buttons for Word/PDF processing
btn_settings_word = tk.Button(frame_right, text="Word/PDF Extraction Settings", command=open_word_extraction_settings)
btn_settings_word.pack(pady=5, fill="x")

btn_process_word = tk.Button(frame_right, text="Process Files (Word/PDF)", command=lambda: process_files("wrd/pdf"))
btn_process_word.pack(pady=5, fill="x")

# Button to export JSON
btn_export_json = tk.Button(root, text="Export to JSON", command=lambda: export_to_json(processed_data))
btn_export_json.pack(pady=20, padx=10, fill="x")

# Progress bar
progress_bar = ttk.Progressbar(root, orient="horizontal", length=100, mode="determinate")
progress_bar.pack(pady=20, padx=10, fill="x")

# Start application
root.mainloop()